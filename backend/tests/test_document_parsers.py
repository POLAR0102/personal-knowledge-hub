from pathlib import Path

import pymupdf
import pytest
from docx import Document as create_docx

from app.ai.document_processing.base import (
    DocumentReadError,
    EmptyDocumentError,
    ScannedPdfNotSupportedError,
    UnsupportedDocumentTypeError,
)
from app.ai.document_processing.registry import get_parser


def test_txt_parser_reads_utf8(tmp_path: Path) -> None:
    file_path = tmp_path / "notes.txt"
    file_path.write_text("第一段\n第二段", encoding="utf-8")

    result = get_parser("txt").parse(file_path)

    assert len(result.sections) == 1
    assert result.sections[0].content.splitlines() == ["第一段", "第二段"]
    assert result.sections[0].section_index == 0


def test_txt_parser_reads_utf8_sig(tmp_path: Path) -> None:
    file_path = tmp_path / "bom.txt"
    file_path.write_text("带 BOM 的内容", encoding="utf-8-sig")

    result = get_parser(".TXT").parse(file_path)

    assert result.sections[0].content == "带 BOM 的内容"


def test_txt_parser_rejects_invalid_encoding(tmp_path: Path) -> None:
    file_path = tmp_path / "invalid.txt"
    file_path.write_bytes(b"\xff\xfe\xfa")

    with pytest.raises(DocumentReadError, match="UTF-8"):
        get_parser("txt").parse(file_path)


def test_markdown_parser_recognizes_headings(tmp_path: Path) -> None:
    file_path = tmp_path / "readme.md"
    file_path.write_text(
        "# 项目介绍\n正文内容\n\n## 技术栈\nFastAPI",
        encoding="utf-8",
    )

    result = get_parser("md").parse(file_path)

    assert [section.section_title for section in result.sections] == [
        "项目介绍",
        "技术栈",
    ]
    assert "# 项目介绍" in result.sections[0].content
    assert "FastAPI" in result.sections[1].content


def test_markdown_parser_preserves_code_blocks(tmp_path: Path) -> None:
    file_path = tmp_path / "code.md"
    file_path.write_text(
        "# 示例\n```python\n# 这不是标题\nprint('hello')\n```",
        encoding="utf-8",
    )

    result = get_parser("md").parse(file_path)

    assert len(result.sections) == 1
    assert "```python" in result.sections[0].content
    assert "# 这不是标题" in result.sections[0].content
    assert "print('hello')" in result.sections[0].content


def test_pdf_parser_extracts_multiple_pages_with_page_numbers(tmp_path: Path) -> None:
    file_path = tmp_path / "pages.pdf"
    with pymupdf.open() as pdf:
        first_page = pdf.new_page()
        first_page.insert_text((72, 72), "First page")
        second_page = pdf.new_page()
        second_page.insert_text((72, 72), "Second page")
        pdf.save(file_path)

    result = get_parser("pdf").parse(file_path)

    assert [section.page_number for section in result.sections] == [1, 2]
    assert [section.section_index for section in result.sections] == [0, 1]
    assert "First page" in result.sections[0].content
    assert "Second page" in result.sections[1].content
    assert result.metadata["page_count"] == 2


def test_pdf_parser_rejects_pdf_without_extractable_text(tmp_path: Path) -> None:
    file_path = tmp_path / "blank.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.save(file_path)

    with pytest.raises(ScannedPdfNotSupportedError, match="OCR"):
        get_parser("pdf").parse(file_path)


def test_docx_parser_extracts_paragraphs(tmp_path: Path) -> None:
    file_path = tmp_path / "paragraphs.docx"
    document = create_docx()
    document.add_paragraph("第一段")
    document.add_paragraph("第二段")
    document.save(file_path)

    result = get_parser("docx").parse(file_path)

    assert result.sections[0].content == "第一段\n第二段"


def test_docx_parser_recognizes_heading_styles(tmp_path: Path) -> None:
    file_path = tmp_path / "headings.docx"
    document = create_docx()
    document.add_heading("章节标题", level=1)
    document.add_paragraph("章节正文")
    document.save(file_path)

    result = get_parser("docx").parse(file_path)

    assert result.sections[0].section_title == "章节标题"
    assert result.sections[0].content == "章节标题\n章节正文"


def test_docx_parser_extracts_table_text_in_order(tmp_path: Path) -> None:
    file_path = tmp_path / "table.docx"
    document = create_docx()
    document.add_paragraph("表格之前")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "角色"
    table.cell(1, 0).text = "张三"
    table.cell(1, 1).text = "负责人"
    document.add_paragraph("表格之后")
    document.save(file_path)

    result = get_parser("docx").parse(file_path)
    content = result.sections[0].content

    assert content.index("表格之前") < content.index("姓名 | 角色")
    assert content.index("姓名 | 角色") < content.index("表格之后")
    assert "张三 | 负责人" in content


@pytest.mark.parametrize("file_type", ["txt", "md", "docx"])
def test_text_parsers_reject_empty_documents(
    tmp_path: Path,
    file_type: str,
) -> None:
    file_path = tmp_path / f"empty.{file_type}"
    if file_type == "docx":
        create_docx().save(file_path)
    else:
        file_path.write_text("  \n\n", encoding="utf-8")

    with pytest.raises(EmptyDocumentError):
        get_parser(file_type).parse(file_path)


def test_registry_rejects_unsupported_document_type() -> None:
    with pytest.raises(UnsupportedDocumentTypeError):
        get_parser("xlsx")
